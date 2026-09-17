#!/usr/bin/env python3
"""Extract the unique identifier-only ortholog table from Supplementary Data 8."""

from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import load_workbook


def main(source: Path, output: Path) -> None:
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook["ABA"]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    if tuple(header[:5]) != ("OGID", "Ath", "Esa", "Sir", "Spa"):
        raise RuntimeError(f"unexpected identifier header: {header[:5]}")
    seen = set()
    records = []
    for row in rows:
        record = tuple("" if value is None else str(value).strip() for value in row[:5])
        if not record[0] or record in seen:
            continue
        seen.add(record)
        records.append(record)
    workbook.close()
    if len(records) != 15198:
        raise RuntimeError(f"expected 15198 unique ortholog rows, got {len(records)}")
    with output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(["author_orthogroup", "ath_gene", "esa_gene", "sir_gene", "spa_gene"])
        writer.writerows(records)
    print(f"rows={len(records)} output={output}")


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
