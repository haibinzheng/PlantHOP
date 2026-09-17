#!/usr/bin/env python3
"""Read-only inventory of GSE268881 supplementary XLSX workbooks."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def clean(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def inspect(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets = []
    for worksheet in workbook.worksheets:
        preview = []
        for row in worksheet.iter_rows(min_row=1, max_row=min(8, worksheet.max_row), values_only=True):
            preview.append([clean(value) for value in row[:30]])
        sheets.append(
            {
                "title": worksheet.title,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "preview": preview,
            }
        )
    workbook.close()
    return {"path": str(path), "bytes": path.stat().st_size, "sheets": sheets}


if __name__ == "__main__":
    results = [inspect(Path(arg).resolve()) for arg in sys.argv[1:]]
    json.dump(results, sys.stdout, ensure_ascii=False, indent=2)
    print()

