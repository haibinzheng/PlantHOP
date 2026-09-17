#!/usr/bin/env python3
"""Expression-blind audit of GSE268881 supplementary mapping tables."""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


FROZEN = {
    "OG0000931": ["AT1G64160", "AT4G11180", "AT4G11190"],
    "OG0001788": ["AT1G20850", "AT4G35350"],
    "OG0001353": ["AT2G28755", "AT2G28760", "AT3G46440"],
    "OG0000478": ["AT5G03170", "AT5G60490"],
    "OG0000555": ["AT1G26570", "AT3G29360", "AT5G15490"],
    "OG0000146": ["AT1G07747", "AT1G48750", "AT1G64235"],
    "OG0001133": ["AT1G68810", "AT3G25710"],
    "OG0001988": ["AT2G14890", "AT2G23130", "AT4G37450"],
    "OG0000560": ["AT1G55330", "AT3G13520", "AT4G26320"],
    "OG0000074": ["AT1G22900", "AT1G55210", "AT1G58170"],
    "OG0005923": ["AT3G27200"],
    "OG0003252": ["AT3G10080"],
    "OG0009497": ["AT5G16490"],
    "OG0002185": ["AT1G78040", "AT4G08685", "AT5G10130"],
    "OG0002797": ["AT3G46300", "AT3G46310", "AT5G02640"],
    "OG0003815": ["AT2G28410", "AT5G60650"],
    "OG0002167": ["AT1G52150", "AT4G32880"],
    "OG0000891": ["AT1G75280", "AT1G75290", "AT1G75300"],
}
TARGETS = {gene for genes in FROZEN.values() for gene in genes}
AT_RE = re.compile(r"AT[1-5CM]G\d{5}", re.IGNORECASE)


def norm(value):
    return "" if value is None else str(value).strip()


def genes_in(value):
    return {m.upper() for m in AT_RE.findall(norm(value))}


def main(folder: Path, output: Path) -> None:
    inventory = []
    hits = []
    keyword_rows = []
    for path in sorted(folder.glob("*.xlsx")):
        if path.stat().st_size == 0:
            continue
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            inv = {
                "workbook": path.name,
                "sheet": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
            }
            inventory.append(inv)
            for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                vals = [norm(v) for v in row]
                joined = " | ".join(vals)
                lower = joined.lower()
                found = sorted(set().union(*(genes_in(v) for v in vals)) & TARGETS)
                if found:
                    hits.append(
                        {
                            "workbook": path.name,
                            "sheet": ws.title,
                            "row": row_idx,
                            "target_genes": found,
                            "values": vals,
                        }
                    )
                if row_idx <= 20 and any(k in lower for k in ("ortholog", "orthogroup", "1-to-1", "one-to-one", "15198")):
                    keyword_rows.append(
                        {
                            "workbook": path.name,
                            "sheet": ws.title,
                            "row": row_idx,
                            "values": vals[:30],
                        }
                    )
        wb.close()

    by_gene = {gene: [] for gene in sorted(TARGETS)}
    for hit in hits:
        for gene in hit["target_genes"]:
            by_gene[gene].append(
                {k: hit[k] for k in ("workbook", "sheet", "row", "values")}
            )

    result = {
        "source_folder": str(folder.resolve()),
        "frozen_orthogroups": FROZEN,
        "inventory": inventory,
        "keyword_rows": keyword_rows,
        "hits": hits,
        "by_gene": by_gene,
    }
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "workbooks": len({x["workbook"] for x in inventory}),
        "sheets": len(inventory),
        "hit_rows": len(hits),
        "target_genes_found": sum(bool(v) for v in by_gene.values()),
        "output": str(output.resolve()),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
